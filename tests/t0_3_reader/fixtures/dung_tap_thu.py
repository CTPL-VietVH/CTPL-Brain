"""Dựng tập thử TỰ DỰNG cho T0.3 ở cả bốn định dạng, từ cùng một nội dung nguồn.

Chạy: .venv/bin/python tests/t0_3_reader/fixtures/dung_tap_thu.py

⚠️ Tập này **không thay thế** 20 văn bản hành chính VN thật mà `docs/08` T0.3
đòi. Xem `tests/t0_3_reader/README.md`.

Điểm quan trọng nhất: bản `.docx` được dựng **CỐ Ý KHÔNG GÁN HEADING STYLE** —
mọi đoạn đều là `Normal`, đánh số bằng tay, in đậm bằng tay. Đó đúng là hình
dạng mà `docs/08` B2 điểm 1 cảnh báo là *phổ biến* trong văn bản hành chính VN
soạn tay, và là lý do bộ chuẩn hoá regex phải là đường chính.
"""

from __future__ import annotations

import pathlib
import sys

HERE = pathlib.Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

from nguon import KHONG_CAU_TRUC, QUY_CHE, QUY_TRINH  # noqa: E402

FONT_VN = "/System/Library/Fonts/Supplemental/Arial Unicode.ttf"


def dung_txt_md() -> None:
    (HERE / "quy_che.txt").write_text(QUY_CHE, encoding="utf-8")
    (HERE / "quy_trinh.txt").write_text(QUY_TRINH, encoding="utf-8")
    (HERE / "khong_cau_truc.txt").write_text(KHONG_CAU_TRUC, encoding="utf-8")

    # Bản markdown: cùng nội dung, nhưng tiêu đề mang dấu '#'. Bộ đọc phải hạ
    # dấu '#' rồi để regex quyết cấu trúc — nếu markdown tự dựng cây theo '#'
    # thì sẽ có hai bộ quy tắc cấu trúc và chúng sẽ lệch nhau.
    md = []
    for ln in QUY_CHE.split("\n"):
        if ln.startswith("CHƯƠNG "):
            md.append(f"## {ln}")
        elif ln.startswith("Điều "):
            md.append(f"### {ln}")
        else:
            md.append(ln)
    (HERE / "quy_che.md").write_text("\n".join(md), encoding="utf-8")
    print("✓ .txt và .md")


def dung_docx() -> None:
    import docx
    from docx.shared import Pt

    d = docx.Document()
    for ln in QUY_CHE.split("\n"):
        p = d.add_paragraph()
        run = p.add_run(ln)
        # ⛔ CỐ Ý: style luôn là Normal. Chương/Điều chỉ được in đậm bằng tay,
        # đúng thói quen soạn thảo thật — KHÔNG có Heading style nào để mà tin.
        p.style = d.styles["Normal"]
        if ln.startswith(("CHƯƠNG ", "Điều ")):
            run.bold = True
        run.font.size = Pt(13)
    d.save(HERE / "quy_che_khong_style.docx")

    # Đối chứng: cùng nội dung NHƯNG có gán Heading style tử tế.
    d2 = docx.Document()
    for ln in QUY_CHE.split("\n"):
        if ln.startswith("CHƯƠNG "):
            d2.add_paragraph(ln, style="Heading 1")
        elif ln.startswith("Điều "):
            d2.add_paragraph(ln, style="Heading 2")
        else:
            d2.add_paragraph(ln)
    d2.save(HERE / "quy_che_co_style.docx")
    print("✓ .docx (một bản không style, một bản có style để đối chứng)")


def dung_pdf() -> None:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfgen import canvas

    pdfmetrics.registerFont(TTFont("VN", FONT_VN))
    c = canvas.Canvas(str(HERE / "quy_che.pdf"), pagesize=A4)
    rong, cao = A4
    le, dong_cao = 56, 15
    y = cao - 56

    for ln in QUY_CHE.split("\n"):
        if y < 56:
            c.showPage()
            y = cao - 56
        c.setFont("VN", 11)
        # Xuống dòng thủ công cho dòng dài
        while len(ln) > 95:
            cat = ln.rfind(" ", 0, 95)
            cat = cat if cat > 0 else 95
            c.drawString(le, y, ln[:cat])
            ln = ln[cat:].lstrip()
            y -= dong_cao
            if y < 56:
                c.showPage()
                y = cao - 56
                c.setFont("VN", 11)
        c.drawString(le, y, ln)
        y -= dong_cao
    c.save()
    print("✓ .pdf (có lớp chữ, phông Unicode)")


if __name__ == "__main__":
    dung_txt_md()
    dung_docx()
    dung_pdf()
    print(f"\nTập thử tự dựng nằm ở: {HERE}")
    print("⚠️  KHÔNG thay thế 20 văn bản thật mà docs/08 T0.3 đòi.")
